import io
import zipfile

import pytest
from PIL import Image
from pptx import Presentation

from brand import blog, social

COPY = {
    "headline_light": {
        "accent": "Aetherion",
        "headline": "now talks to Jira",
        "body": "Connect Jira once and every agent can use it.",
    },
    "headline_dark": {
        "eyebrow": "Product update",
        "headline": "Your agents can read Jira",
        "body": "Connect it once from the Apps page.",
    },
    "three_steps": {
        "accent": "Jira",
        "headline": "in three steps",
        "step1_title": "Connect",
        "step1_text": "Open Apps and click Connect.",
        "step2_title": "Approve",
        "step2_text": "Sign in to Jira and approve.",
        "step3_title": "Use",
        "step3_text": "Every agent can now read Jira.",
    },
}


def _slide_text(pptx: bytes) -> str:
    prs = Presentation(io.BytesIO(pptx))
    assert len(prs.slides) == 1
    texts = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                walk(sh.shapes)
            elif sh.has_text_frame:
                texts.append(sh.text_frame.text)

    walk(prs.slides[0].shapes)
    return "\n".join(texts)


@pytest.mark.parametrize("layout", list(COPY))
def test_render_fills_png_and_pptx(layout):
    out = social.render(layout, COPY[layout])
    img = Image.open(io.BytesIO(out["png"]))
    assert img.size == (1080, 1350)
    text = _slide_text(out["pptx"])
    for value in COPY[layout].values():
        assert value in text or value.upper() in text
    assert "Innovation" not in text  # template copy is fully replaced
    with zipfile.ZipFile(io.BytesIO(out["pptx"])) as z:
        assert any(n.startswith("ppt/fonts/") for n in z.namelist())  # brand fonts stay embedded


def test_copy_that_does_not_fit_is_rejected_with_a_limit():
    too_long = {**COPY["headline_light"], "accent": "An accent far too long for the slot"}
    with pytest.raises(ValueError, match="accent .*limit about 16"):
        social.render("headline_light", too_long)


def test_slots_must_match_the_layout():
    with pytest.raises(ValueError, match="missing \\['body'\\]"):
        social.validate_copy("headline_light", {"accent": "A", "headline": "B"})
    with pytest.raises(ValueError, match="Unknown layout"):
        social.validate_copy("poster", {})


def test_upper_slots_are_uppercased():
    clean = social.validate_copy("headline_dark", COPY["headline_dark"])
    assert clean["eyebrow"] == "PRODUCT UPDATE"


def test_layout_guide_lists_limits():
    guide = social.layout_guide()
    assert guide["three_steps"]["slots"]["step1_text"]["max_chars"] == 80


BLOG = """# Your agents can read Jira

Teams track work in **Jira**. Now agents can too.

## What's new
- Connect once
1. Use everywhere

Closing *line*.
"""


def test_blog_parse():
    title, blocks = blog.parse(BLOG)
    assert title == "Your agents can read Jira"
    assert [k for k, _ in blocks] == ["p", "h2", "bullet", "bullet", "p"]
    assert blog._runs("a **b** *c*") == [
        ("a ", False, False),
        ("b", True, False),
        (" ", False, False),
        ("c", False, True),
    ]


def test_blog_documents_build():
    out = blog.render(BLOG)
    assert out["docx"][:2] == b"PK" and out["pdf"][:4] == b"%PDF"
    from docx import Document

    doc = Document(io.BytesIO(out["docx"]))
    assert doc.paragraphs[0].text == "Your agents can read Jira"
