"""Rebuild src/brand/assets from the Calfus social-media template (dev-time, macOS).

    .venv/bin/python scripts/build_brand_assets.py "Calfus -Social media templates.pptx"

Writes, for each layout in src/brand/layouts.json:
  - assets/layouts/<layout>.pptx  the template cut down to that one slide (fonts stay embedded)
  - assets/plates/<layout>.png    the slide with every text run blanked, rendered by Quick Look
and extracts the template's embedded fonts (OFL) into assets/fonts/.
Run it again whenever marketing updates the template, then check layouts.json shape ids.
"""

import json
import re
import struct
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation

BRAND = Path(__file__).resolve().parents[1] / "src" / "brand"
ASSETS = BRAND / "assets"
SIZE = (1080, 1350)

# Embedded font full name -> file name the renderer loads.
FONTS = {
    "DM Sans 1 Normal": "DMSans-Regular.ttf",
    "DM Sans 1 SemiBold": "DMSans-SemiBold.ttf",
    "DM Sans 1 Bold": "DMSans-Bold.ttf",
    "DM Sans 1 Black": "DMSans-Black.ttf",
    "Stack Sans Notch Light": "StackSansNotch-Light.ttf",
    "Stack Sans Notch Regular": "StackSansNotch-Regular.ttf",
    "Stack Sans Notch SemiBold": "StackSansNotch-SemiBold.ttf",
    "Stack Sans Notch Bold": "StackSansNotch-Bold.ttf",
}


def _eot_font(data: bytes) -> tuple[str, bytes] | None:
    """Full name and TrueType bytes of an uncompressed EOT (.fntdata); None if compressed."""
    eot_size, font_size, _version, flags = struct.unpack_from("<IIII", data, 0)
    if flags & 0x4:  # MTX-compressed subset: not usable as a TTF
        return None
    off = 16 + 10 + 1 + 1 + 4 + 2 + 2 + 16 + 8 + 4 + 16
    names = []
    for _ in range(4):  # family, style, version, full name
        off += 2
        (n,) = struct.unpack_from("<H", data, off)
        off += 2
        names.append(data[off : off + n].decode("utf-16-le"))
        off += n
    font = data[eot_size - font_size : eot_size]
    if flags & 0x10000000:
        font = bytes(b ^ 0x50 for b in font)
    return names[3], font


def extract_fonts(template: Path) -> None:
    out = ASSETS / "fonts"
    out.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(template) as z:
        for name in z.namelist():
            if name.startswith("ppt/fonts/") and name.endswith(".fntdata"):
                found = _eot_font(z.read(name))
                if found and found[0] in FONTS:
                    (out / FONTS[found[0]]).write_bytes(found[1])
    missing = [f for f in FONTS.values() if not (out / f).is_file()]
    if missing:
        raise SystemExit(f"Template no longer embeds: {missing}")


def single_slide(template: Path, slide_no: int, dest: Path) -> None:
    """Keep one slide; python-pptx's save drops every part nothing references any more."""
    with zipfile.ZipFile(template) as z:
        files = {n: z.read(n) for n in z.namelist()}
    pres = files["ppt/presentation.xml"].decode()
    rels = files["ppt/_rels/presentation.xml.rels"].decode()
    slide_ids = re.findall(r'<p:sldId [^>]*r:id="(rId\d+)"[^>]*/>', pres)
    keep = slide_ids[slide_no - 1]
    for rid in slide_ids:
        if rid != keep:
            pres = re.sub(rf'<p:sldId [^>]*r:id="{rid}"[^>]*/>', "", pres)
            rels = re.sub(rf'<Relationship [^>]*Id="{rid}"[^>]*/>', "", rels)
    files["ppt/presentation.xml"] = pres.encode()
    files["ppt/_rels/presentation.xml.rels"] = rels.encode()
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "raw.pptx"
        with zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED) as z:
            for n, data in files.items():
                z.writestr(n, data)
        Presentation(raw).save(dest)


def plate(layout_pptx: Path, dest: Path) -> None:
    """Render the slide with all text blanked; the renderer draws text on top."""
    prs = Presentation(layout_pptx)

    def blank(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                blank(sh.shapes)
            elif sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        run.text = ""

    blank(prs.slides[0].shapes)
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "plate.pptx"
        prs.save(src)
        subprocess.run(
            ["qlmanage", "-t", "-s", "1350", "-o", tmp, str(src)], check=True, capture_output=True
        )
        # Quick Look rounds the width (1076 px); the renderer's coordinates are 1080x1350.
        Image.open(Path(tmp) / "plate.pptx.png").convert("RGB").resize(SIZE, Image.LANCZOS).save(
            dest, optimize=True
        )


def extract_logo(template: Path, slide_no: int = 1, shape_id: int = 21) -> None:
    """The black Calfus logo is a picture fill on slide 1; the blog documents use it."""
    slide = Presentation(template).slides[slide_no - 1]
    shape = next(s for s in slide.shapes if s.shape_id == shape_id)
    blip = shape._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    (ASSETS / "logo_black.png").write_bytes(slide.part.related_part(rid).blob)


def main(template: Path) -> None:
    layouts = json.loads((BRAND / "layouts.json").read_text())
    extract_fonts(template)
    extract_logo(template)
    for sub in ("layouts", "plates"):
        (ASSETS / sub).mkdir(parents=True, exist_ok=True)
    for name, spec in layouts.items():
        pptx_path = ASSETS / "layouts" / f"{name}.pptx"
        single_slide(template, spec["slide"], pptx_path)
        plate(pptx_path, ASSETS / "plates" / f"{name}.png")
        print(f"{name}: slide {spec['slide']} -> {pptx_path.name}, plate")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(Path(sys.argv[1]))
