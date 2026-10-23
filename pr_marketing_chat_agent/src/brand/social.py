"""Fill a Calfus social-post layout: an editable .pptx and a ready-to-post 1080x1350 .png.

Both come from the same fitted copy: the font size the PNG needs to fit a slot is written into
the .pptx run too, so the two files match. Layouts and their slots live in layouts.json.
"""

import io
import json
from functools import cache
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).parent
ASSETS = ROOT / "assets"
PX_PER_PT = 4 / 3  # the template is 1080 px wide at 96 dpi
MIN_SCALE = 0.6  # never shrink a slot below 60% of the template size
LINE_GAP = 1.15


@cache
def layouts() -> dict[str, Any]:
    return json.loads((ROOT / "layouts.json").read_text(encoding="utf-8"))


def layout_guide() -> dict[str, Any]:
    """What the writer needs per layout: description and slot limits."""
    return {
        name: {
            "description": spec["description"],
            "slots": {k: {"max_chars": s["max_chars"]} for k, s in spec["slots"].items()},
        }
        for name, spec in layouts().items()
    }


@cache
def _font(name: str, px: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(ASSETS / "fonts" / name), px)


def _width(text: str, font: ImageFont.FreeTypeFont, tracking: float = 0) -> float:
    return font.getlength(text) + tracking * max(len(text) - 1, 0)


def _wrap(text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str] | None:
    """Greedy word wrap; None if a single word is wider than the box."""
    lines: list[str] = []
    for paragraph in text.split("\n"):
        line = ""
        for word in paragraph.split():
            trial = f"{line} {word}".strip()
            if _width(trial, font) <= width:
                line = trial
            elif not line or _width(word, font) > width:
                return None
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines


def fit(text: str, slot: dict[str, Any]) -> tuple[float, list[str]]:
    """Largest point size (template size down to MIN_SCALE) at which `text` fits the slot."""
    _, _, w, h = slot["box"]
    base = slot["size_pt"]
    scale = 1.0
    while scale >= MIN_SCALE - 1e-9:
        pt = base * scale
        px = round(pt * PX_PER_PT)
        lines = _wrap(text, _font(slot["font"], px), w)
        if lines and len(lines) <= slot["max_lines"] and len(lines) * px * LINE_GAP <= h * 1.05:
            return pt, lines
        scale -= 0.04
    raise ValueError(f"too long to fit; shorten it (limit about {slot['max_chars']} characters)")


def validate_copy(layout: str, copy: dict[str, str]) -> dict[str, str]:
    """Check the copy against the layout's slots; returns cleaned copy or raises ValueError."""
    spec = layouts().get(layout)
    if not spec:
        raise ValueError(f"Unknown layout {layout!r}. Available: {', '.join(layouts())}")
    slots = spec["slots"]
    unknown = sorted(set(copy) - set(slots))
    missing = sorted(set(slots) - {k for k, v in copy.items() if str(v).strip()})
    if unknown or missing:
        raise ValueError(
            f"Layout {layout} needs exactly these slots: {sorted(slots)}"
            + (f"; missing {missing}" if missing else "")
            + (f"; unknown {unknown}" if unknown else "")
        )
    clean: dict[str, str] = {}
    problems = []
    for name, slot in slots.items():
        text = " ".join(str(copy[name]).split())
        if slot.get("upper"):
            text = text.upper()
        try:
            fit(text, slot)
        except ValueError as err:
            problems.append(f"{name} ({len(text)} chars): {err}")
        clean[name] = text
    if problems:
        raise ValueError("Copy does not fit the layout: " + "; ".join(problems))
    return clean


def _draw_block(
    draw: ImageDraw.ImageDraw,
    lines: list[str],
    box: list[int],
    font: ImageFont.FreeTypeFont,
    color: str,
    align: str = "left",
    valign: str = "top",
    tracking: float = 0,
) -> None:
    x, y, w, h = box
    line_h = font.size * LINE_GAP
    top = y + (h - line_h * len(lines)) / 2 if valign == "middle" else y
    for i, line in enumerate(lines):
        lw = _width(line, font, tracking)
        lx = x + (w - lw) / 2 if align == "center" else x
        ly = top + i * line_h
        if tracking:
            for ch in line:
                draw.text((lx, ly), ch, font=font, fill=f"#{color}")
                lx += font.getlength(ch) + tracking
        else:
            draw.text((lx, ly), line, font=font, fill=f"#{color}")


def render_png(layout: str, copy: dict[str, str]) -> bytes:
    spec = layouts()[layout]
    img = Image.open(ASSETS / "plates" / f"{layout}.png").convert("RGB")
    draw = ImageDraw.Draw(img)
    for item in spec.get("static", []):
        font = _font(item["font"], round(item["size_pt"] * PX_PER_PT))
        _draw_block(
            draw,
            [item["text"]],
            item["box"],
            font,
            item["color"],
            item.get("align", "left"),
            item.get("valign", "top"),
            item.get("tracking", 0),
        )
    for name, slot in spec["slots"].items():
        pt, lines = fit(copy[name], slot)
        font = _font(slot["font"], round(pt * PX_PER_PT))
        _draw_block(
            draw,
            lines,
            slot["box"],
            font,
            slot["color"],
            slot.get("align", "left"),
            slot.get("valign", "top"),
        )
    out = io.BytesIO()
    img.save(out, "PNG", optimize=True)
    return out.getvalue()


def _shape(shapes: Any, shape_id: int) -> Any:
    for sh in shapes:
        if sh.shape_id == shape_id:
            return sh
        if sh.shape_type == 6:  # group
            found = _shape(sh.shapes, shape_id)
            if found is not None:
                return found
    return None


def _set_text(shape: Any, text: str, pt: float, run_index: int | None) -> None:
    """Replace a slot's text, keeping the template's run formatting."""
    paragraphs = shape.text_frame.paragraphs
    runs = [r for p in paragraphs for r in p.runs]
    if run_index is not None:
        runs[run_index].text = text
        runs[run_index].font.size = Pt(pt)
        return
    first = paragraphs[0]
    first.runs[0].text = text
    first.runs[0].font.size = Pt(pt)
    for run in first.runs[1:]:
        run._r.getparent().remove(run._r)
    for br in first._p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}br"):
        first._p.remove(br)
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def render_pptx(layout: str, copy: dict[str, str]) -> bytes:
    spec = layouts()[layout]
    prs = Presentation(ASSETS / "layouts" / f"{layout}.pptx")
    slide = prs.slides[0]
    for name, slot in spec["slots"].items():
        pt, _ = fit(copy[name], slot)
        shape = _shape(slide.shapes, slot["shape"])
        if shape is None:
            raise RuntimeError(f"{layout}: shape {slot['shape']} not in template; rebuild assets")
        _set_text(shape, copy[name], pt, slot.get("run"))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def render(layout: str, copy: dict[str, str]) -> dict[str, bytes]:
    """Validated copy -> {"png": bytes, "pptx": bytes}."""
    clean = validate_copy(layout, copy)
    return {"png": render_png(layout, clean), "pptx": render_pptx(layout, clean)}
